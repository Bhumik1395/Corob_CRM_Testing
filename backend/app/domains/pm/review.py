import io
import re

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from pydantic import BaseModel

from app.core.audit import audit
from app.core.auth import CurrentUser, require_roles
from app.core.rate_limit import rate_limit
from app.domains.pm.outcome import (next_question, count_matches, list_remarks, export_rows,
                                    rows_for_remark, tickets_for_remark, DONE, DISPUTE)
from app.domains.pm.overrides_store import set_overrides, clear_overrides
from app.domains.pm.rules_store import (get_learned_rules, save_rule, list_history, set_label, set_active)
from app.domains.pm.cache import load_pm_df
from app.domains.pm.label_log import log_label
from app.shared.employee_mapping_service import load_employee_mapping, load_supervisor_mapping

router = APIRouter()
employee_only = require_roles("corob_employee", "admin")


class Answer(BaseModel):
    phrase: str
    label: str


class NewLabel(BaseModel):
    label: str


class TicketRef(BaseModel):
    company: str
    ticket_no: str


class RowMark(BaseModel):
    company: str | None = None
    remark_key: str
    label: str | None = None          # None = clear the override (back to the remark's rule)
    tickets: list[TicketRef] = []
    all_rows: bool = False            # True = every row with this remark, across all pages


def _check_label(label):
    if label not in (DONE, DISPUTE):
        raise HTTPException(400, {"error": "bad_label", "message": "Label must be 'PM Done' or 'Dispute'."})


@router.get("/api/pm/review/next")
def review_next(company: str | None = None, skip: str = "", user: CurrentUser = Depends(employee_only)):
    df = load_pm_df(user, company)
    if df is None:
        raise HTTPException(404, {"error": "no_data", "message": "No PM data uploaded yet for this company."})
    return next_question(df, get_learned_rules(), skip=[s for s in skip.split("|") if s])


@router.get("/api/pm/review/remarks")
def review_remarks(
    company: str | None = None,
    status: str = "all",
    search: str = "",
    page: int = 1,
    page_size: int = 25,
    user: CurrentUser = Depends(employee_only),
):
    df = load_pm_df(user, company)
    if df is None:
        raise HTTPException(404, {"error": "no_data", "message": "No PM data uploaded yet for this company."})
    return list_remarks(df, get_learned_rules(), status=status, search=search, page=max(page, 1), page_size=min(max(page_size, 1), 100))


_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_HEADERS = ["Ticket", "Dealer", "Employee name", "Supervisor name", "Date", "Remark"]
_WIDTHS = [24, 36, 26, 26, 13, 80]


def _text(v) -> str:
    return _ILLEGAL.sub("", "" if v is None else str(v))


@router.get("/api/pm/review/export", dependencies=[Depends(rate_limit("export", 20, 60))])
def review_export(
    company: str | None = None,
    status: str = "all",
    search: str = "",
    user: CurrentUser = Depends(employee_only),
):
    if status not in ("all", "unresolved", "done", "dispute"):
        raise HTTPException(400, {"error": "bad_status", "message": "Unknown status filter."})
    df = load_pm_df(user, company)
    if df is None:
        raise HTTPException(404, {"error": "no_data", "message": "No PM data uploaded yet for this company."})

    rows = export_rows(df, get_learned_rules(), status=status, search=search)

    names, sups = load_employee_mapping(), load_supervisor_mapping()
    wb = Workbook()
    ws = wb.active
    ws.title = "Remarks"
    ws.append(_HEADERS)

    top_wrap = Alignment(vertical="top", wrap_text=True)
    n = 0
    for r in rows.to_dict("records"):
        code = str(r.get("eng_code", "") or "").strip()
        emp = names.get(code) or r.get("employee_name") or ""
        sup = r.get("supervisor") or sups.get(code) or ""
        d_name, d_code = str(r.get("dealer_name") or ""), str(r.get("dealer_code") or "")
        dealer = f"{d_name} ({d_code})" if d_name and d_code else d_name or d_code
        d = r.get("call_date")
        date = d.date() if pd.notna(d) and hasattr(d, "date") else None

        ws.append([_text(r.get("ticket_no")), _text(dealer), _text(emp), _text(sup), date, _text(r.get("remarks"))])
        n += 1
        for c in ws[n + 1]:
            if isinstance(c.value, str):
                c.data_type = "s"          # a remark starting with "=" must never become a formula
            c.alignment = top_wrap
        ws.cell(row=n + 1, column=5).number_format = "yyyy-mm-dd"

    for i, w in enumerate(_WIDTHS):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"

    if n:
        table = Table(displayName="PMRemarks", ref=f"A1:F{n + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
        ws.add_table(table)

    buf = io.BytesIO()
    wb.save(buf)
    audit(user.username, "pm_remarks_export", f"{status} / {n} rows")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="PM_Remarks.xlsx"'},
    )

@router.post("/api/pm/review/answer")
def review_answer(body: Answer, user: CurrentUser = Depends(employee_only)):
    _check_label(body.label)
    try:
        save_rule(body.phrase, body.label, user.username)
    except ValueError:
        raise HTTPException(400, {"error": "bad_phrase", "message": "Phrase must be 1 to 255 characters."})
    log_label("phrase_answer", user.username, body.label, phrase=body.phrase)
    return {"saved": True}


@router.get("/api/pm/review/history")
def review_history(company: str | None = None, scope: str = "mine", user: CurrentUser = Depends(employee_only)):
    rules = list_history(None if scope == "all" else user.username)
    counts = {}
    if company:
        df = load_pm_df(user, company)
        if df is not None:
            counts = count_matches(df, [r["phrase"] for r in rules])
    for r in rules:
        r["rowsAffected"] = counts.get(r["phrase"]) if counts else None
    return {"rules": rules}


@router.post("/api/pm/review/rules/{rule_id}/label")
def review_change(rule_id: int, body: NewLabel, user: CurrentUser = Depends(employee_only)):
    _check_label(body.label)
    if not set_label(rule_id, body.label, user.username):
        raise HTTPException(404, {"error": "not_found", "message": "Rule not found."})
    log_label("rule_change", user.username, body.label, rule_id=rule_id)
    return {"changed": True}


@router.post("/api/pm/review/rules/{rule_id}/restore")
def review_restore(rule_id: int, user: CurrentUser = Depends(employee_only)):
    if not set_active(rule_id, True, user.username):
        raise HTTPException(404, {"error": "not_found", "message": "Rule not found."})
    log_label("rule_restore", user.username, None, rule_id=rule_id)
    return {"restored": True}


@router.delete("/api/pm/review/rules/{rule_id}")
def review_remove(rule_id: int, user: CurrentUser = Depends(employee_only)):
    if not set_active(rule_id, False, user.username):
        raise HTTPException(404, {"error": "not_found", "message": "Rule not found."})
    log_label("rule_remove", user.username, None, rule_id=rule_id)
    return {"removed": True}


@router.get("/api/pm/review/rows")
def review_rows(
    key: str,
    company: str | None = None,
    page: int = 1,
    page_size: int = 25,
    user: CurrentUser = Depends(employee_only),
):
    df = load_pm_df(user, company)
    if df is None:
        raise HTTPException(404, {"error": "no_data", "message": "No PM data uploaded yet for this company."})
    return rows_for_remark(df, get_learned_rules(), key, max(page, 1), min(max(page_size, 1), 100))


@router.post("/api/pm/review/rows/mark")
def review_rows_mark(body: RowMark, user: CurrentUser = Depends(employee_only)):
    if body.label is not None:
        _check_label(body.label)
    if body.all_rows:
        df = load_pm_df(user, body.company)
        if df is None:
            raise HTTPException(404, {"error": "no_data", "message": "No PM data uploaded yet for this company."})
        pairs = tickets_for_remark(df, body.remark_key)
    else:
        if len(body.tickets) > 5000:
            raise HTTPException(400, {"error": "too_many", "message": "Select at most 5000 rows at once."})
        pairs = [(t.company, t.ticket_no.strip()) for t in body.tickets]
    if not pairs:
        raise HTTPException(400, {"error": "nothing_selected", "message": "No rows selected."})
    if body.label is None:
        clear_overrides(pairs)
    else:
        set_overrides(pairs, body.label, user.username)
    log_label(
        "row_reset" if body.label is None else "row_mark",
        user.username, body.label,
        remark_key=body.remark_key,
        company=body.company,
        all_rows=body.all_rows,
        ticket_count=len(pairs),
        tickets=[{"company": c, "ticket_no": t} for c, t in pairs[:5000]],
    )
    return {"updated": len(pairs)}